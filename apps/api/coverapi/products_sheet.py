"""The workshop's product list linked to the drawing covers (ADR-091).

The owner, 7 October 2026: "we have an Excel sheet in which all products are linked to the
drawings we uploaded; please link them to the right drawings so we can sort easily" — "I'd like
to upload it" — "it will be a zip".

At the Desk a person uploads the list: an .xlsx or .csv, or a .zip holding one or more of them
(and perhaps drawing PDFs). The studio's drawings-zip button hands a zip with a sheet and no
drawings here too. The program

- reads every sheet, finds each sheet's header row (the first row with two or more text cells);
- finds the column that links a row to a drawing: the one whose values best match the drawing
  covers' codes ("S40", "C 23", "s40", "cover 39 & 46 - L1 & L5 mirror") or the order number on
  the drawing ("Cover 105", "COVER-20"), tolerant of case, spaces, a "cover" prefix and leading
  zeros. A person may name another column and apply again;
- picks by content, not by header, the description (the longest text) and a short number
  column ("Cover 1"): the owner's list heads its description column with a fabric colour;
- links the SUNS catalogue models a description names: its families ("Portofino/ Aspen/ Kota")
  and its type ("D-Bed", "lounge chair", "2-seater", "Left"). Family and type agree: linked.
  Family only, with no type to compare: a suggestion, which a person confirms by eye;
- writes per drawing cover `products.json` and per linked SUNS model `drawings.json`, which the
  Desk shows, searches and sorts by;
- keeps the uploaded file in the data folder's products/ (a history), and reports what was
  linked, what was not, which covers have no product and which codes carry two products.

A zip is unpacked in memory with care: no absolute paths or "..", no more than
`products.max_files` entries and `products.max_bytes` unpacked; __MACOSX and hidden files are
skipped. PDFs in it are only listed and matched to existing drawing covers by their code; no
model is made from them.
"""

from __future__ import annotations

import csv
import io
import json
import re
import time
import zipfile
from collections.abc import Callable
from pathlib import Path, PurePosixPath
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile

SHEETS = (".xlsx", ".xlsm", ".csv")
PREVIEW_ROWS = 8  # param-ok: rows shown back to the person
LABEL_CHARS = 80  # param-ok: a product label on the Desk at most this long
NUMBER_CHARS = 20  # param-ok: a "number" cell ("Cover 105") is at most this long
SHOWN = 200  # param-ok: unmatched rows sent back at most
PRODUCTS_JSON = "products.json"
DRAWINGS_JSON = "drawings.json"
CODE_TOKEN = re.compile(r"[a-z]\d+[a-z]?")
Keys = dict[str, list[str]]


class SheetError(ValueError):
    """Something the person can fix: said plainly."""


# ---- matching codes -------------------------------------------------------------------------


def norm(value: Any) -> str:
    """A code made comparable: lower case, only letters and digits, no leading zeros in a
    number ("Cover 0105" and "cover105" are one; "C 23" and "c23" too)."""
    s = re.sub(r"[^a-z0-9]", "", str(value or "").lower())
    return re.sub(r"(?<![0-9])0+(?=[0-9])", "", s)


def _order_number(pdf: Path) -> str | None:
    """The drawing's own order number ("Order Number: Cover 105", "COVER-20")."""
    try:
        import pymupdf

        text = " ".join(p.get_text() for p in pymupdf.open(pdf).pages())
    except Exception:  # noqa: BLE001 - a drawing that cannot be read has no order number
        return None
    m = re.search(r"order\s*number\s*:?\s*(cover[\s_-]*\d+[a-z]?)", text, re.I)
    if not m:
        m = re.search(r"\b(cover[\s_-]*\d+[a-z]?)\b", text, re.I)
    return m.group(1) if m else None


def _cached_order(d: Path, data_dir: Path, drawing_pdf: Callable[..., Path | None]) -> str | None:
    """The PDF's order number, read once and kept in the cover's folder."""
    cache = d / "order_number.txt"
    if cache.is_file():
        return cache.read_text(encoding="utf-8").strip() or None
    pdf = drawing_pdf(d, data_dir)
    order = _order_number(pdf) if pdf is not None else None
    try:
        cache.write_text(order or "", encoding="utf-8")
    except OSError:
        pass
    return order


def drawing_keys(models: Path, data_dir: Path) -> Keys:
    """Every way a drawing cover may be named in a sheet → its model ids.

    First pass: the cover's whole code, its folder name and the order number on its PDF. Second
    pass, for what is still free: the codes inside a combined name ("cover 39 & 46 - L1 & L5
    mirror" answers to L5, not to L1, which drawing-l1 has) and variants ("S10 box" and "S10
    plain" both answer to S10)."""
    from coverapi.desk import _code, drawing_pdf

    dirs = [d for d in sorted(models.glob("drawing-*")) if (d / "cover.json").is_file()]
    keys: Keys = {}
    for d in dirs:
        names = {_code(d), d.name.removeprefix("drawing-")}
        order = _cached_order(d, data_dir, drawing_pdf)
        if order:
            names.add(order)
        for k in {norm(n) for n in names}:
            if k and d.name not in keys.setdefault(k, []):
                keys[k].append(d.name)
    second: Keys = {}
    for d in dirs:
        text = f"{_code(d)} {d.name.removeprefix('drawing-')}".lower()
        for tok in re.split(r"[^a-z0-9]+", text):
            k = norm(tok)
            if CODE_TOKEN.fullmatch(tok) and k not in keys and d.name not in second.get(k, []):
                second.setdefault(k, []).append(d.name)
    return keys | second


def match(value: Any, keys: Keys) -> list[str]:
    """The drawing covers a cell names: the whole value, else each part ("S40 / S41")."""
    k = norm(value)
    if not k:
        return []
    if k in keys:
        return keys[k]
    out: list[str] = []
    for part in re.split(r"[,;/+|&]|\s{2,}", str(value)):
        for mid in keys.get(norm(part), []):
            if mid not in out:
                out.append(mid)
    return out


# ---- reading sheets -------------------------------------------------------------------------


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _table(rows: list[list[Any]], name: str) -> dict[str, Any] | None:
    """One sheet as {name, header, rows}: the header is the first row with two text cells."""
    clean = [[_cell(v) for v in r] for r in rows]
    for i, r in enumerate(clean):
        texts = [c for c in r if c and not re.fullmatch(r"[\d.,\s-]+", c)]
        if len(texts) >= 2:  # param-ok: a header names at least two columns
            header = [c or f"column {j + 1}" for j, c in enumerate(r)]
            seen: dict[str, int] = {}
            for j, h in enumerate(header):  # two columns with one name stay apart
                seen[h] = seen.get(h, 0) + 1
                if seen[h] > 1:
                    header[j] = f"{h} ({seen[h]})"
            body = [dict(zip(header, rr + [""] * (len(header) - len(rr)), strict=False))
                    for rr in clean[i + 1:] if any(rr)]  # fmt: skip
            return {"name": name, "header": header, "rows": body}
    return None


def read_sheets(filename: str, data: bytes) -> list[dict[str, Any]]:
    """The tables in one .xlsx or .csv file."""
    low = filename.lower()
    if low.endswith(".xls"):
        raise SheetError(f"{filename}: an old .xls file; please save it as .xlsx (Excel: Save As,"
                         " Excel Workbook) and upload again")  # fmt: skip
    if low.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect: Any = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        t = _table([list(r) for r in csv.reader(io.StringIO(text), dialect)], filename)
        return [t] if t else []
    if low.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - say it plainly
            raise SheetError(f"{filename}: not a readable Excel file ({exc})") from None
        out = []
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            t = _table(rows, f"{filename} · {ws.title}")
            if t:
                out.append(t)
        return out
    raise SheetError(f"{filename}: upload an .xlsx, a .csv or a .zip")


def unpack(
    data: bytes, max_files: int, max_bytes: int
) -> tuple[list[tuple[str, bytes]], list[str]]:
    """The sheets and the PDFs in a zip, safely. Refused as a whole when an entry's path
    leaves the zip ("..", an absolute path) or when it is too large."""
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise SheetError("that is not a readable zip file") from None
    infos = [i for i in z.infolist() if not i.is_dir()]
    if len(infos) > max_files:
        raise SheetError(f"the zip holds {len(infos)} files; at most {max_files}")
    for info in infos:  # every path first: one bad entry refuses the whole zip
        name = info.filename.replace("\\", "/")
        drive = re.match(r"^[a-zA-Z]:", name)
        if name.startswith("/") or ".." in PurePosixPath(name).parts or drive:
            raise SheetError(f"the zip holds an unsafe path ({info.filename!r}); nothing was read")
    total = 0
    sheets: list[tuple[str, bytes]] = []
    pdfs: list[str] = []
    for info in infos:
        p = PurePosixPath(info.filename.replace("\\", "/"))
        if "__MACOSX" in p.parts or any(part.startswith(".") for part in p.parts):
            continue
        total += info.file_size
        if total > max_bytes:
            raise SheetError(f"the zip unpacks to more than {max_bytes // 1_000_000} MB")
        low = p.name.lower()
        if low.endswith((*SHEETS, ".xls")):
            with z.open(info) as f:
                raw = f.read(info.file_size + 1)
            if len(raw) != info.file_size:  # a lying header: refuse
                raise SheetError(f"{p.name}: its size in the zip is not true; nothing was read")
            sheets.append((p.name, raw))
        elif low.endswith(".pdf"):
            pdfs.append(p.name)
    return sheets, pdfs


def zip_has_sheet(path: Path) -> bool:
    """For the drawings-zip button: does this zip hold a product list?"""
    try:
        with zipfile.ZipFile(path) as z:
            return any(PurePosixPath(n).name.lower().endswith((*SHEETS, ".xls"))
                       and "__MACOSX" not in n for n in z.namelist())  # fmt: skip
    except zipfile.BadZipFile:
        return False


# ---- which column is what -------------------------------------------------------------------


def best_column(tables: list[dict[str, Any]], keys: Keys) -> str | None:
    """The column whose values name drawing covers most often (over every table)."""
    hits: dict[str, int] = {}
    for t in tables:
        for col in t["header"]:
            hits[col] = hits.get(col, 0) + sum(1 for r in t["rows"] if match(r.get(col), keys))
    best = max(hits.items(), key=lambda kv: kv[1], default=(None, 0))
    return best[0] if best[1] > 0 else None


def roles(table: dict[str, Any], link_col: str) -> tuple[str | None, str | None]:
    """The description column (longest text with letters) and a number column (short cells
    with a digit, like "Cover 1"), read from the cells; the link column is neither."""
    desc, num = None, None
    best_len, best_num = 0.0, 0.0
    for col in table["header"]:
        if col == link_col:
            continue
        vals = [str(r.get(col) or "") for r in table["rows"] if r.get(col)]
        if not vals:
            continue
        letters = sum(1 for v in vals if re.search(r"[A-Za-z]", v)) / len(vals)
        avg = sum(len(v) for v in vals) / len(vals)
        if letters > 0.5 and avg > best_len:  # param-ok: most cells hold words
            desc, best_len = col, avg
    for col in table["header"]:
        if col in (link_col, desc):
            continue
        vals = [str(r.get(col) or "") for r in table["rows"] if r.get(col)]
        if not vals:
            continue
        short = sum(1 for v in vals if len(v) <= NUMBER_CHARS and re.search(r"\d", v)) / len(vals)
        if short > 0.5 and short > best_num:  # param-ok: most cells are short numbers
            num, best_num = col, short
    return desc, num


# ---- SUNS families and types ----------------------------------------------------------------

# words in SUNS names that say what a piece is, not which family it belongs to
GENERIC = set(
    "suns lounge lounger loungetable chair chairs seater seat sofa bench corner table tabel side "
    "dining bar low big small single middle hocker daybed chaise longue sunlounger left right "
    "rigth part with without arm arms open angled curved bended rounded moon moonshape organic "
    "point extension collection high back weaving cross fishbone macrame knot alu aluminium teak "
    "hpl neolith grc dia large medium relax picnic fire pit dropshape roof shape set normal and "
    "the for lounge cover covers ibiza style".split()
)


def kinds(text: str) -> tuple[set[str], set[str]]:
    """What a description or a SUNS name says the piece is: its type and its side."""
    t = re.sub(r"\bw/o\b.*", "", text.lower())  # "w/o side table": not a table
    t = t.replace("_", " ").replace("-", " ")
    for word, digit in (("two", "2"), ("three", "3"), ("four", "4")):
        t = re.sub(rf"\b{word}\b", digit, t)
    types: set[str] = set()
    if re.search(r"\bd\s*bed|daybed|lounger|sun\s*lounger", t):
        types.add("daybed")
    if re.search(r"chaise|\blongue\b|\bcl\b", t):
        types.add("chaise")
    if re.search(r"(dining|bar)\s+chair|bar\s*chair", t):
        types.add("dining-chair")
    elif re.search(r"lounge\s+chair|single\s+seat|\bchair\b", t):
        types.add("chair")
    for m in re.finditer(r"\b(\d)(?:\s*[,. ]\s*(5))?\s*sea?r?ter", t):
        types.add("seater-" + m.group(1) + ("-5" if m.group(2) else ""))
    if re.search(r"\bbench\b", t):
        types.add("bench")
    if re.search(r"corr?ner|\bhoek\b|\bl\s*part\b", t):
        types.add("corner")
    if re.search(r"\bside\s+table|\btable\s+side", t):
        types.add("side-table")
    elif re.search(r"(dining|bar)\s*table|bartable|\bdining\b", t):
        types.add("dining-table")
    elif re.search(r"\btable|\btabel\b|loungetable", t):
        types.add("table")
    if re.search(r"\bhocker\b|footstool", t):
        types.add("hocker")
    if re.search(r"\bmiddle\b", t):
        types.add("middle")
    if re.search(r"\bmoon", t):
        types.add("moon")
    sides = {s for s in ("left", "right") if re.search(rf"\b{s}\b", t)}
    if re.search(r"\brigth\b", t):
        sides.add("right")
    if re.search(r"\bangle", t):  # an angled piece is a shape of its own: never "the" 2-seater
        sides.add("angled")
    if re.search(r"\bmoon", t):
        sides.add("moon")
    return types, sides


TABLES = {"table", "side-table", "dining-table", "dining-chair"}  # not part of a lounge set


def _words(text: str) -> set[str]:
    return {w for w in re.split(r"[^a-z]+", text.lower()) if len(w) > 2}  # param-ok: no "cl"


def suns_index(models: Path) -> dict[str, dict[str, Any]]:
    """Per SUNS model: its family words and its type, from its name."""
    out = {}
    for d in sorted(models.glob("suns-*")):
        if (d / "cover.json").is_file():
            name = d.name.removeprefix("suns-")
            types, sides = kinds(name)
            out[d.name] = {"families": _words(name) - GENERIC, "types": types, "sides": sides}
    return out


def suns_for(description: str, index: dict[str, dict[str, Any]]) -> dict[str, list[str]]:
    """The SUNS models a description names. Linked: a family and a type in common, the shape
    (angled, moon) the same and the side not crossed. Suggested: a family in common where the
    description names no type (a "lounge set"), tables only for a table."""
    vocab = set().union(*(m["families"] for m in index.values())) if index else set()
    fams = _words(description) & vocab
    types, sides = kinds(description)
    shape = sides - {"left", "right"}
    linked, suggested = [], []
    for mid, m in index.items():
        if not fams & m["families"]:
            continue
        lr, mlr = sides & {"left", "right"}, m["sides"] & {"left", "right"}
        if lr and mlr and not lr & mlr:
            continue  # a left piece is not the right one
        mshape = m["sides"] - {"left", "right"}
        if types & m["types"] and shape == mshape:
            linked.append(mid)
        elif not types and not (m["types"] & TABLES) and shape <= mshape:
            suggested.append(mid)
    return {"linked": linked, "suggested": suggested}


# ---- linking --------------------------------------------------------------------------------


def link(tables: list[dict[str, Any]], keys: Keys, column: str | None,
         index: dict[str, dict[str, Any]] | None = None) -> dict[str, Any]:  # fmt: skip
    """Every row to its drawing covers through `column` (or the best column), and to the SUNS
    models its description names."""
    col = column or best_column(tables, keys)
    if col is None:
        raise SheetError("no column names the drawings (S40, C23, Cover 105 …): choose the column")
    if not any(col in t["header"] for t in tables):
        raise SheetError(f"there is no column {col!r}")
    linked: dict[str, list[dict[str, Any]]] = {}
    unmatched: list[dict[str, Any]] = []
    for t in tables:
        if col not in t["header"]:
            continue
        desc, num = roles(t, col)
        for r in t["rows"]:
            label = [v[:LABEL_CHARS] for v in (r.get(desc or ""), r.get(num or "")) if v]
            row = {**r, "_sheet": t["name"], "_labels": label,
                   "_suns": suns_for(r.get(desc or "", ""), index or {})}  # fmt: skip
            mids = match(r.get(col), keys)
            for mid in mids:
                linked.setdefault(mid, []).append(row)
            if not mids and any(v for k, v in r.items() if k != col):
                unmatched.append(row)
    by_code: dict[str, set[str]] = {}  # one code on two different products: worth a look
    for rows in linked.values():
        for r in rows:
            if r["_labels"]:
                by_code.setdefault(str(r.get(col)).strip(), set()).add(r["_labels"][0])
    twice = {c: sorted(v) for c, v in sorted(by_code.items()) if len(v) > 1}
    return {"column": col, "linked": linked, "unmatched": unmatched, "twice": twice}


def _write_json(path: Path, doc: dict[str, Any]) -> None:
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(doc, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    tmp.replace(path)


def write(models: Path, result: dict[str, Any], source: str, by: str) -> dict[str, Any]:
    """Each drawing cover's products.json and each SUNS model's drawings.json; covers and
    models no longer in the list lose theirs. Returns the SUNS counts."""
    from coverapi.desk import _code

    col, now = result["column"], time.time()
    suns: dict[str, dict[str, list[dict[str, str]]]] = {}
    for d in sorted(models.glob("drawing-*")):
        rows = result["linked"].get(d.name)
        path = d / PRODUCTS_JSON
        if not rows:
            path.unlink(missing_ok=True)
            continue
        lk = sorted({m for r in rows for m in r["_suns"]["linked"]})
        sg = sorted({m for r in rows for m in r["_suns"]["suggested"]} - set(lk))
        _write_json(path, {
            "source": source, "column": col, "by": by, "time": now,
            "rows": [{k: v for k, v in r.items() if k not in ("_labels", "_suns")} for r in rows],
            "labels": sorted({lab for r in rows for lab in r["_labels"]}),
            "suns": {"linked": lk, "suggested": sg},
        })  # fmt: skip
        desc = next((r["_labels"][0] for r in rows if r["_labels"]), "")
        ref = {"id": d.name, "code": _code(d), "product": desc}
        for kind, ids in (("linked", lk), ("suggested", sg)):
            for m in ids:
                suns.setdefault(m, {"linked": [], "suggested": []})[kind].append(ref)
    for d in sorted(models.glob("suns-*")):
        path = d / DRAWINGS_JSON
        if d.name in suns:
            _write_json(path, {"source": source, "by": by, "time": now, **suns[d.name]})
        else:
            path.unlink(missing_ok=True)
    return {
        "linked": sum(1 for v in suns.values() if v["linked"]),
        "suggested_only": sum(1 for v in suns.values() if not v["linked"]),
    }


def _doc(path: Path) -> dict[str, Any]:
    try:
        doc = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return doc if isinstance(doc, dict) else {}


def product_labels(d: Path) -> list[str]:
    """What the Desk lists, searches and sorts by: a drawing cover's products, or the drawings
    a SUNS model is linked to ("D1 · Portofino/ Aspen/ Kota normal D-Bed")."""
    if d.name.startswith("suns-"):
        return [f"{r['code']} · {r['product']}".strip(" ·")
                for r in _doc(d / DRAWINGS_JSON).get("linked") or []]  # fmt: skip
    return [str(x) for x in _doc(d / PRODUCTS_JSON).get("labels") or []]


def products_card(d: Path) -> dict[str, Any]:
    """The card's Products part: the sheet's rows, and the SUNS links either way."""
    if d.name.startswith("suns-"):
        doc = _doc(d / DRAWINGS_JSON)
        return {"drawings": {"linked": doc.get("linked") or [],
                             "suggested": doc.get("suggested") or []}} if doc else {}  # fmt: skip
    doc = _doc(d / PRODUCTS_JSON)
    if not doc:
        return {}
    rows = doc.get("rows") or []
    columns = [c for c in (rows[0] if rows else {}) if not c.startswith("_")]
    return {"source": doc.get("source"), "column": doc.get("column"), "time": doc.get("time"),
            "columns": columns, "rows": [[r.get(c, "") for c in columns] for r in rows],
            "suns": doc.get("suns") or {"linked": [], "suggested": []}}  # fmt: skip


# ---- the endpoint ---------------------------------------------------------------------------


def install(app: FastAPI, store: Any) -> Callable[..., dict[str, Any]]:
    """The Desk's product-list upload; returns `importer(name, data, by, column)` for the
    drawings-zip button, which hands over a zip that holds a sheet and no drawings."""
    from coverengine.params import Registry

    from coverapi.desk import approver
    from coverapi.security import current_user

    folder = store.root / "products"

    def allowed(request: Request) -> Any:
        user = current_user(request)
        if not (user.may("edit") or approver(user, Registry.load(None).resolve())):
            raise HTTPException(403, "only editors and the Desk's approvers")
        return user

    def process(name: str, data: bytes, column: str | None, by: str) -> dict[str, Any]:
        p = Registry.load(None).resolve()
        files: list[tuple[str, bytes]]
        pdfs: list[str] = []
        if name.lower().endswith(".zip"):
            files, pdfs = unpack(data, int(p["products.max_files"]),  # type: ignore[arg-type]
                                 int(p["products.max_bytes"]))  # type: ignore[arg-type]  # fmt: skip
            if not files and not pdfs:
                raise SheetError("the zip holds no .xlsx, .csv or .pdf")
        else:
            files = [(name, data)]
        per_file: list[dict[str, Any]] = []
        tables: list[dict[str, Any]] = []
        for fname, raw in files:
            try:
                ts = read_sheets(fname, raw)
            except SheetError as exc:
                per_file.append({"file": fname, "error": str(exc)})
                continue
            per_file.append({"file": fname, "tables": len(ts),
                             "rows": sum(len(t["rows"]) for t in ts)})  # fmt: skip
            tables += ts
        if files and not tables:
            errors = [f["error"] for f in per_file if "error" in f]
            raise SheetError("; ".join(errors) or "the list holds no table with a header row")
        models = Path(store.models)
        keys = drawing_keys(models, store.root)
        out: dict[str, Any] = {"files": per_file}
        if tables:
            res = link(tables, keys, column, suns_index(models))
            suns = write(models, res, name, by)
            drawings = sorted(d.name for d in models.glob("drawing-*")
                              if (d / "cover.json").is_file())  # fmt: skip
            clean = [
                {k: v for k, v in r.items() if not k.startswith("_")} for r in res["unmatched"]
            ]
            out |= {
                "column": res["column"],
                "columns": list(dict.fromkeys(c for t in tables for c in t["header"])),
                "rows": sum(len(t["rows"]) for t in tables),
                "linked_rows": len({id(r) for v in res["linked"].values() for r in v}),
                "covers": len(res["linked"]),
                "unmatched": clean[:SHOWN],
                "unmatched_count": len(clean),
                "without_products": [d for d in drawings if d not in res["linked"]],
                "twice": res["twice"],
                "suns": suns,
                "preview": [
                    {"covers": match(r.get(res["column"]), keys), **r}
                    for t in tables
                    for r in t["rows"]
                ][:PREVIEW_ROWS],  # fmt: skip
            }
        if pdfs:
            # "Cover 105 - S40.pdf": the code after the order number, else the whole name
            found = {
                f: match(re.sub(r"^cover\s*\d+\w?[^-]*-\s*", "", Path(f).stem, flags=re.I), keys)
                or match(Path(f).stem, keys)
                for f in pdfs
            }
            out["pdfs"] = {"count": len(pdfs), "matched": sum(1 for v in found.values() if v),
                           "new": sorted(k for k, v in found.items() if not v)}  # fmt: skip
        return out

    def keep(name: str, data: bytes) -> None:
        folder.mkdir(parents=True, exist_ok=True)
        stamp = time.strftime("%Y%m%d-%H%M%S")
        (folder / f"sheet-{stamp}-{re.sub(r'[^A-Za-z0-9._-]', '_', name)}").write_bytes(data)

    def importer(name: str, data: bytes, by: str, column: str | None = None) -> dict[str, Any]:
        keep(name, data)
        try:
            out = process(name, data, column, by)
        except SheetError as exc:
            raise HTTPException(400, str(exc)) from None
        return {**out, "file": name, "kind": "products"}

    @app.post("/api/desk-products")
    async def upload_products(
        request: Request,
        file: UploadFile | None = File(None),  # noqa: B008
        column: str = Form(""),  # noqa: B008
    ) -> dict[str, Any]:
        """The product list (.xlsx, .csv or a .zip of them) linked to the drawing covers. Without
        a file: the last upload applied again (with another `column`)."""
        user = allowed(request)
        by = user.name or user.username
        limit = int(Registry.load(None).resolve()["products.max_bytes"])  # type: ignore[arg-type]
        if file is not None:
            name = Path(file.filename or "products.xlsx").name
            data = await file.read(limit + 1)
            if len(data) > limit:
                raise HTTPException(400, f"the file is larger than {limit // 1_000_000} MB")
            return importer(name, data, by, column.strip() or None)
        last = sorted(folder.glob("sheet-*")) if folder.is_dir() else []
        if not last:
            raise HTTPException(400, "upload the product list first")
        name = last[-1].name.split("-", 3)[-1]
        try:
            out = process(name, last[-1].read_bytes(), column.strip() or None, by)
        except SheetError as exc:
            raise HTTPException(400, str(exc)) from None
        return {**out, "file": name, "kind": "products"}

    @app.get("/api/desk-products")
    def last_products(request: Request) -> dict[str, Any]:
        """Which product list was uploaded last, and when."""
        current_user(request)
        last = sorted(folder.glob("sheet-*")) if folder.is_dir() else []
        if not last:
            return {"file": None}
        return {"file": last[-1].name.split("-", 3)[-1], "time": last[-1].stat().st_mtime,
                "uploads": len(last)}  # fmt: skip

    return importer
